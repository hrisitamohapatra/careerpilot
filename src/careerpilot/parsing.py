"""Resume parsing (PDF / DOCX / pasted text) and section-aware chunking."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

from . import config
from .security import ValidationError, clean_text, validate_upload

SECTION_HEADINGS = {
    "Summary": ["summary", "profile", "objective", "about me", "about"],
    "Skills": ["skills", "technical skills", "core skills", "tech stack"],
    "Experience": ["experience", "work experience", "internships", "employment", "internship"],
    "Projects": ["projects", "project work", "personal projects"],
    "Education": ["education", "academics", "academic background"],
    "Certifications": ["certifications", "certificates", "achievements", "awards"],
}
_HEADING_LOOKUP = {alias: name for name, aliases in SECTION_HEADINGS.items() for alias in aliases}


@dataclass
class Chunk:
    text: str
    section: str


def extract_text(filename: str, data: bytes) -> str:
    """Validate then extract text from an uploaded PDF or DOCX (in memory)."""
    ext = validate_upload(filename, data)
    try:
        if ext == ".pdf":
            import fitz  # pymupdf

            with fitz.open(stream=data, filetype="pdf") as doc:
                pages = [doc[i].get_text() for i in range(min(len(doc), config.MAX_PDF_PAGES))]
            text = "\n".join(pages)
        else:
            from docx import Document

            d = Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
    except ValidationError:
        raise
    except Exception:
        raise ValidationError("Could not read this file. Try another file or paste the text.")
    text = clean_text(text)
    if len(text) < 30:
        raise ValidationError("No readable text found (scanned PDF?). Paste your resume text instead.")
    return text


def _heading_of(line: str) -> str | None:
    key = re.sub(r"[^a-z ]", "", line.lower()).strip()
    if key and len(line) <= 40:
        return _HEADING_LOOKUP.get(key)
    return None


def split_sections(text: str, max_chars: int = 2000) -> list[Chunk]:
    """Split by heading heuristics; fall back to generic ~500-token chunks."""
    sections: list[Chunk] = []
    current, buf = "Summary", []
    found = False
    for line in text.splitlines():
        h = _heading_of(line.strip())
        if h:
            found = True
            if "".join(buf).strip():
                sections.append(Chunk("\n".join(buf).strip(), current))
            current, buf = h, []
        else:
            buf.append(line)
    if "".join(buf).strip():
        sections.append(Chunk("\n".join(buf).strip(), current))
    if not found:
        sections = [Chunk(text, "General")]
    out: list[Chunk] = []
    for s in sections:  # ~500 tokens is about 2000 chars
        t = s.text
        while len(t) > max_chars:
            cut = t.rfind("\n", 0, max_chars)
            cut = cut if cut > 200 else max_chars
            out.append(Chunk(t[:cut].strip(), s.section))
            t = t[cut:]
        if t.strip():
            out.append(Chunk(t.strip(), s.section))
    return out
